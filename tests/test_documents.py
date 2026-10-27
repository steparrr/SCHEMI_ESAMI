import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "scripts"))
from common import create_exam, load_config
from pipeline import Pipeline
from ingest_helpers import chunk_blocks


class DocumentTests(unittest.TestCase):
    def test_long_single_paragraph_is_split_without_losing_words(self):
        text = " ".join(f"parola{i}" for i in range(5000))
        chunks = chunk_blocks([{"type": "text", "text": text, "heading": "Alfa", "seq": 0}], "DOC", "fonte.txt")
        self.assertGreaterEqual(len(chunks), 3)
        self.assertEqual(" ".join(chunk["text"] for chunk in chunks).split(), text.split())

    def test_docx_images_are_extracted_and_sent_with_correct_id(self):
        with tempfile.TemporaryDirectory() as temporary:
            base = Path(temporary)
            # Il motore usa la configurazione reale; soltanto i dati sono temporanei.
            root = create_exam("Esame immagine", base)
            (root / "input" / "programma.txt").write_text("Alfa | alfa\n")
            png = base / "immagine.png"
            image = Image.new("RGB", (600, 400), "white")
            draw = ImageDraw.Draw(image)
            draw.rectangle((80, 80, 480, 300), fill="#8BBBD9")
            draw.text((100, 120), "ALFA - IMMAGINE TECNICA", fill="black")
            image.save(png)
            document = Document()
            document.add_heading("Concetto alfa", 1)
            document.add_paragraph("Alfa: la figura mostra un esempio tecnico. " * 40)
            document.add_picture(str(png))
            document.add_paragraph("Alfa: osservare il rettangolo e la sua posizione nella figura.")
            document.save(root / "input" / "sbobine" / "appunti.docx")
            cfg = load_config(root, PROJECT)
            Pipeline(root, cfg).tick()
            pack = json.loads((root / "source_packs" / "schema_01" / "source_pack.json").read_text())
            self.assertEqual(len(pack["images"]), 1)
            picture = pack["images"][0]
            self.assertTrue((root / picture["path"]).is_file())
            batch = next((root / "batch").glob("*.jsonl"))
            content = json.loads(batch.read_text())["body"]["input"][0]["content"]
            visual = [item for item in content if item["type"] == "input_image"]
            self.assertEqual(len(visual), 1)
            self.assertTrue(visual[0]["image_url"].startswith("data:image/png;base64,"))
            self.assertIn(picture["image_id"], content[-2]["text"])


if __name__ == "__main__":
    unittest.main()
