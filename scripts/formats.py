from __future__ import annotations

import csv
import io
import tempfile
from pathlib import Path

import fitz
from PIL import Image, ImageOps

import ingest_helpers as h
from ocr import ocr_image

IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".heic", ".heif", ".tif", ".tiff"}


def normalized_image(source):
    if isinstance(source, Path) and source.suffix.lower() in {".heic", ".heif"}:
        from pillow_heif import register_heif_opener
        register_heif_opener()
    with Image.open(source) as image:
        output = io.BytesIO()
        ImageOps.exif_transpose(image).convert("RGB").save(output, format="PNG")
        return output.getvalue()


def image_text(data):
    with tempfile.TemporaryDirectory(prefix="schemi-ocr-") as directory:
        path = Path(directory) / "pagina.png"
        path.write_bytes(data)
        return h.clean_text(ocr_image(path))


def extract_pdf(path):
    blocks, media = [], {}
    with fitz.open(path) as document:
        if document.needs_pass:
            raise ValueError("Il PDF è protetto da password: carica una copia sbloccata.")
        for index, page in enumerate(document, 1):
            text = h.clean_text(page.get_text("text"))
            largest = 0
            for picture in page.get_images(full=True):
                try:
                    record = document.extract_image(picture[0])
                    # Normalizzazione: i dati inviati all'API devono essere PNG/JPEG.
                    data = normalized_image(io.BytesIO(record["image"]))
                    media[f"pagina{index}_xref{picture[0]}.png"] = (data, index, None)
                    largest = max([largest] + [r.get_area() / max(1, page.rect.get_area()) for r in page.get_image_rects(picture[0])])
                except (ValueError, OSError):
                    # Il raster dell'intera pagina recupera anche figure non decodificabili.
                    data = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes("png")
                    media[f"pagina{index}_completa.png"] = (data, index, None)
            if len(text) < 20 or (len(text) < 180 and largest > .6):
                raster = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
                recognized = image_text(raster)
                if recognized:
                    text = recognized if len(recognized) >= len(text) else text + "\n" + recognized
                else:
                    media[f"pagina{index}_visiva.png"] = (raster, index, None)
                    text = (text + "\n[Pagina iconografica senza testo riconoscibile; interpretare soltanto la figura fornita.]").strip()
            heading = next((line[:180] for line in text.splitlines() if line.strip()), f"Pagina {index}")
            blocks.append({"type": "page", "text": text, "heading": heading, "seq": index - 1, "page_start": index, "page_end": index})
        return blocks, media, document.page_count


def spreadsheet_blocks(path):
    blocks = []
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="", errors="strict") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            for index, row in enumerate(csv.reader(stream, dialect), 1):
                text = " | ".join(str(value).strip() for value in row)
                if text.strip(" |"):
                    blocks.append({"type": "table_row", "heading": path.stem, "text": f"Riga {index}: {text}", "seq": len(blocks)})
        return blocks
    import openpyxl
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in formulas.worksheets:
            values_sheet = cached[worksheet.title]
            for row in worksheet.iter_rows():
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        computed = values_sheet.cell(cell.row, cell.column).value
                        value = f"{value} [risultato: {computed if computed is not None else 'non disponibile nel file'}]"
                    values.append("" if value is None else str(value))
                text = " | ".join(values)
                if text.strip(" |"):
                    blocks.append({"type": "table_row", "heading": worksheet.title, "text": f"{worksheet.title}, riga {row[0].row}: {text}", "seq": len(blocks)})
    finally:
        formulas.close()
        cached.close()
    return blocks


def extract_pptx(path):
    from pptx import Presentation
    blocks, media = [], {}
    presentation = Presentation(path)

    def shapes_in(shapes):
        for shape in shapes:
            yield shape
            if hasattr(shape, "shapes"):
                yield from shapes_in(shape.shapes)

    for number, slide in enumerate(presentation.slides, 1):
        parts = []
        heading = slide.shapes.title.text if slide.shapes.title is not None else f"Slide {number}"
        for shape in shapes_in(slide.shapes):
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text)
            if shape.has_table:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            if shape.has_chart:
                for series in shape.chart.series:
                    parts.append(f"Grafico — {series.name}: " + ", ".join(str(value) for value in series.values))
            if hasattr(shape, "image"):
                data = normalized_image(io.BytesIO(shape.image.blob))
                media[f"slide{number}_{shape.shape_id}.png"] = (data, number, number - 1)
                if not parts or len(" ".join(parts)) < 80:
                    recognized = image_text(data)
                    if recognized:
                        parts.append(recognized)
        # Recupera anche le etichette testuali di diagrammi non esposti come textbox.
        known = "\n".join(parts)
        for node in slide.element.iter():
            if node.tag.endswith("}t") and node.text and node.text not in known:
                parts.append(node.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                parts.append("Note del docente: " + notes)
        text = h.clean_text("\n".join(parts)) or "[Slide iconografica senza testo; interpretare soltanto le figure fornite.]"
        blocks.append({"type": "slide", "heading": heading, "text": text, "seq": number - 1, "page_start": number, "page_end": number})
    return blocks, media, len(presentation.slides)
