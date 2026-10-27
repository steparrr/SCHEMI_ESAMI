from __future__ import annotations
import csv, hashlib, io, json, re, shutil, sqlite3, subprocess, sys, unicodedata, zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
import xml.etree.ElementTree as ET
import fitz
from PIL import Image, ImageStat
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()

def clean_text(s: str) -> str:
    if not s:
        return ""
    s = s.replace("\u00ad","").replace("\ufeff","")
    s = re.sub(r"([A-Za-zÀ-ÖØ-öø-ÿ])\-\s*\n\s*([a-zà-öø-ÿ])", r"\1\2", s)
    s = s.replace("\r\n","\n").replace("\r","\n")
    lines, blank = [], False
    for line in s.split("\n"):
        line = re.sub(r"[ \t]+"," ",line).strip()
        if not line:
            if not blank:
                lines.append("")
            blank = True
        else:
            lines.append(line)
            blank = False
    return "\n".join(lines).strip()

def words(s: str):
    return re.findall(r"\b[\wÀ-ÖØ-öø-ÿ’'-]+\b", s, flags=re.UNICODE)

def token_estimate(s: str) -> int:
    return max(1, int(round(len(words(s)) * 1.35))) if s.strip() else 0

def normalize_question_key(s: str) -> str:
    s = unicodedata.normalize("NFKD", s.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

def next_numeric_id(rows, key: str, prefix: str, width: int) -> int:
    best = 0
    rx = re.compile(rf"^{re.escape(prefix)}(\d+)$")
    for r in rows:
        m = rx.match(str(r.get(key) or ""))
        if m:
            best = max(best, int(m.group(1)))
    return best + 1

def create_empty_db(path: Path):
    con = sqlite3.connect(path)
    con.executescript("""
    PRAGMA journal_mode=WAL;
    CREATE TABLE IF NOT EXISTS documents(document_id TEXT PRIMARY KEY, filename TEXT, original_path TEXT, format TEXT, source_type TEXT, size_bytes INTEGER, sha256 TEXT, pages INTEGER, has_extractable_text INTEGER, has_images INTEGER, status TEXT, processed_at TEXT);
    CREATE TABLE IF NOT EXISTS chunks(chunk_id TEXT PRIMARY KEY, document_id TEXT, document_name TEXT, page_start INTEGER, page_end INTEGER, heading TEXT, text TEXT, token_count INTEGER, content_hash TEXT, sequence_number INTEGER);
    CREATE TABLE IF NOT EXISTS questions(question_id TEXT PRIMARY KEY, question_text TEXT, answer_text TEXT, source_document TEXT, source_document_id TEXT, location TEXT, year TEXT, professor TEXT, frequency INTEGER, topic TEXT, content_hash TEXT, extraction_mode TEXT);
    CREATE TABLE IF NOT EXISTS images(image_id TEXT, document_id TEXT, source_document TEXT, page_or_position TEXT, path TEXT, filename_originale TEXT, width INTEGER, height INTEGER, format TEXT, sha256 TEXT, perceptual_hash TEXT, caption TEXT, heading TEXT, context_before TEXT, context_after TEXT, status TEXT, filtered_reason TEXT);
    CREATE TABLE IF NOT EXISTS schemes(numero INTEGER PRIMARY KEY, titolo TEXT, descrizione TEXT, keywords TEXT, stato TEXT);
    CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(chunk_id UNINDEXED, heading, text, document_name, tokenize='unicode61 remove_diacritics 2');
    CREATE VIRTUAL TABLE IF NOT EXISTS questions_fts USING fts5(question_id UNINDEXED, question_text, professor, topic, tokenize='unicode61 remove_diacritics 2');
    """)
    con.commit()
    con.close()

def dhash(img: Image.Image, hash_size=8):
    im = img.convert("L").resize((hash_size+1,hash_size))
    px = list(im.getdata()); bits = []
    for y in range(hash_size):
        row = px[y*(hash_size+1):(y+1)*(hash_size+1)]
        bits.extend(row[x] > row[x+1] for x in range(hash_size))
    v = 0
    for b in bits:
        v = (v << 1) | int(b)
    return f"{v:0{hash_size*hash_size//4}x}"

def inspect_image(data: bytes):
    try:
        im = Image.open(io.BytesIO(data)); im.load()
        w,h = im.size; fmt = (im.format or "").lower()
        gray = im.convert("L").resize((min(w,256),min(h,256)))
        stat = ImageStat.Stat(gray)
        mean,std = stat.mean[0],stat.stddev[0]
        reason = None
        if w < 80 and h < 80:
            reason = "too_small"
        elif mean > 248 and std < 4:
            reason = "near_blank"
        return w,h,fmt,dhash(im),reason
    except Exception:
        return None,None,None,None,"unreadable"

def docx_pages(path: Path):
    try:
        with zipfile.ZipFile(path) as z:
            if "docProps/app.xml" not in z.namelist():
                return None
            root = ET.fromstring(z.read("docProps/app.xml"))
            for el in root.iter():
                if el.tag.endswith("Pages") and el.text and el.text.isdigit():
                    return int(el.text)
    except Exception:
        pass
    return None

def docx_extract(path: Path):
    doc = Document(str(path))
    blocks=[]; current_heading=""; seq=0
    body=doc.element.body
    for child in body.iterchildren():
        tag=child.tag.split("}")[-1]
        if tag=="p":
            p=Paragraph(child,doc)
            txt=clean_text(p.text)
            style=(p.style.name if p.style else "") or ""
            if (style.lower().startswith("heading") or style.lower().startswith("titolo")) and txt:
                current_heading=txt
            rel_ids=[]
            for el in child.iter():
                if el.tag.endswith("}blip"):
                    rid=el.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if rid: rel_ids.append(rid)
            if txt or rel_ids:
                blocks.append({"type":"paragraph","text":txt,"heading":current_heading,"style":style,"seq":seq,"rel_ids":rel_ids})
                seq+=1
        elif tag=="tbl":
            tbl=Table(child,doc)
            rows=[]; rel_ids=[]
            for row in tbl.rows:
                vals=[]
                for cell in row.cells:
                    vals.append(clean_text(cell.text).replace("\n"," | "))
                    for p in cell.paragraphs:
                        for el in p._p.iter():
                            if el.tag.endswith("}blip"):
                                rid=el.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                                if rid: rel_ids.append(rid)
                rows.append("\t".join(vals))
            txt=clean_text("\n".join(rows))
            if txt or rel_ids:
                blocks.append({"type":"table","text":txt,"heading":current_heading,"style":"Table","seq":seq,"rel_ids":rel_ids})
                seq+=1
    relmap={}
    for rid,rel in doc.part.rels.items():
        try:
            if "image" in rel.reltype:
                relmap[rid]=rel.target_ref
        except Exception:
            pass
    media={}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("word/media/") and not n.endswith("/"):
                media[n.split("/")[-1]]=z.read(n)
    locations=defaultdict(list)
    for b in blocks:
        for rid in b["rel_ids"]:
            tgt=relmap.get(rid)
            if tgt:
                locations[Path(tgt).name].append({"seq":b["seq"],"heading":b["heading"]})
    return blocks,media,locations

def pdf_extract(path: Path):
    doc=fitz.open(path)
    blocks=[]; page_texts=[]
    for i,p in enumerate(doc):
        txt=clean_text(p.get_text("text"))
        if len(txt)<20 and shutil.which("tesseract"):
            pix=p.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
            tmp=STATE_DIR/f"_ocr_{path.stem}_{i+1}.png"
            pix.save(str(tmp))
            try:
                proc=subprocess.run(["tesseract",str(tmp),"stdout","-l","ita+eng","--psm","6"],capture_output=True,text=True,timeout=90)
                txt=clean_text(proc.stdout)
            except Exception:
                pass
            tmp.unlink(missing_ok=True)
        page_texts.append(txt)
        ls=[x.strip() for x in txt.split("\n") if x.strip()]
        heading=ls[0][:180] if ls else f"Pagina {i+1}"
        blocks.append({"type":"page","text":txt,"heading":heading,"seq":i,"page_start":i+1,"page_end":i+1})
    return blocks,page_texts,doc.page_count

def ocr_image(path: Path):
    if not shutil.which("tesseract"):
        raise RuntimeError("Tesseract non disponibile per OCR immagine.")
    proc=subprocess.run(["tesseract",str(path),"stdout","-l","ita+eng","--psm","6"],capture_output=True,text=True,timeout=120)
    return clean_text(proc.stdout)

def image_context(blocks, seq, radius=2):
    if seq is None:
        return "",""
    before=[]; after=[]
    for b in blocks:
        if b["seq"]<seq and b.get("text"): before.append(b["text"])
        if b["seq"]>seq and b.get("text"): after.append(b["text"])
    return clean_text("\n".join(before[-radius:]))[-1200:], clean_text("\n".join(after[:radius]))[:1200]

def chunk_blocks(blocks, document_id, filename, min_words=500, max_words=1500):
    chunks=[]; buf=[]; nwords=0; cur_heading=""; page_start=None; page_end=None
    def flush():
        nonlocal buf,nwords,cur_heading,page_start,page_end
        if not buf: return
        text=clean_text("\n\n".join(x["text"] for x in buf if x.get("text")))
        if text:
            chunks.append({
                "document_id":document_id,"document_name":filename,
                "page_start":page_start,"page_end":page_end,
                "heading":cur_heading or (buf[0].get("heading") or ""),
                "text":text,"sequence_number":len(chunks)+1,
                "token_count":token_estimate(text),
                "content_hash":sha256_bytes(text.encode())
            })
        buf=[]; nwords=0; cur_heading=""; page_start=None; page_end=None
    for b in blocks:
        txt=b.get("text","")
        if not txt: continue
        wc=len(words(txt)); h=b.get("heading") or ""
        if buf and h and h!=cur_heading and nwords>=min_words:
            flush()
        if not buf:
            cur_heading=h; page_start=b.get("page_start"); page_end=b.get("page_end")
        if b.get("page_start"):
            page_start=page_start or b["page_start"]; page_end=b.get("page_end") or b["page_start"]
        if wc>max_words:
            if buf: flush()
            paras=[]
            for paragraph in re.split(r"\n\n+",txt):
                if not paragraph.strip():
                    continue
                tokens=paragraph.split()
                if len(tokens)>max_words:
                    paras.extend(" ".join(tokens[i:i+max_words]) for i in range(0,len(tokens),max_words))
                else:
                    paras.append(paragraph)
            temp=[]; tw=0
            for ptxt in paras:
                pw=len(words(ptxt))
                if temp and tw+pw>max_words:
                    t=clean_text("\n\n".join(temp))
                    chunks.append({"document_id":document_id,"document_name":filename,"page_start":b.get("page_start"),"page_end":b.get("page_end"),"heading":h,"text":t,"sequence_number":len(chunks)+1,"token_count":token_estimate(t),"content_hash":sha256_bytes(t.encode())})
                    temp=[]; tw=0
                temp.append(ptxt); tw+=pw
            if temp:
                t=clean_text("\n\n".join(temp))
                chunks.append({"document_id":document_id,"document_name":filename,"page_start":b.get("page_start"),"page_end":b.get("page_end"),"heading":h,"text":t,"sequence_number":len(chunks)+1,"token_count":token_estimate(t),"content_hash":sha256_bytes(t.encode())})
            continue
        if buf and nwords+wc>max_words:
            flush(); cur_heading=h; page_start=b.get("page_start"); page_end=b.get("page_end")
        buf.append(b); nwords+=wc
    flush()
    if len(chunks)>=2 and len(words(chunks[-1]["text"]))<200 and len(words(chunks[-2]["text"]))+len(words(chunks[-1]["text"]))<=1700:
        last=chunks.pop(); prev=chunks[-1]
        prev["text"]=clean_text(prev["text"]+"\n\n"+last["text"])
        prev["page_end"]=last["page_end"] or prev["page_end"]
        prev["token_count"]=token_estimate(prev["text"])
        prev["content_hash"]=sha256_bytes(prev["text"].encode())
    for i,c in enumerate(chunks,1):
        c["sequence_number"]=i
    return chunks

def extract_question_units(text: str, mode="generic"):
    lines=[x.strip() for x in text.replace("\u200b","").split("\n") if x.strip()]
    if mode=="pdf":
        units=[]; cur=[]
        for line in lines:
            if re.match(r"^(?:x\s+)?APPELLO|^\d+\s+appello$|^PRIMO DI APPELLO|^SECONDO APPELLO",line,re.I):
                continue
            if line=="-":
                if cur: units.append(" ".join(cur).strip()); cur=[]
                continue
            if line.startswith("-"):
                if cur: units.append(" ".join(cur).strip()); cur=[]
                cur=[line.lstrip("-").strip()]
            else:
                cur.append(line)
        if cur: units.append(" ".join(cur).strip())
        return [clean_text(u) for u in units if len(words(u))>=2]
    return [clean_text(x) for x in lines if len(words(x))>=2]

def spreadsheet_question_units(path: Path):
    ext=path.suffix.lower()
    out=[]
    if ext==".csv":
        with path.open("r",encoding="utf-8-sig",errors="replace",newline="") as f:
            for row in csv.reader(f):
                s=" | ".join(clean_text(x) for x in row if clean_text(x))
                if len(words(s))>=2: out.append(s)
        return out
    try:
        import openpyxl
    except Exception:
        raise RuntimeError("openpyxl non installato per file XLSX.")
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals=[clean_text(str(x)) for x in row if x is not None and clean_text(str(x))]
            s=" | ".join(vals)
            if len(words(s))>=2: out.append(s)
    return out

def rebuild_fts(con):
    con.execute("DELETE FROM chunks_fts")
    con.execute("INSERT INTO chunks_fts(chunk_id,heading,text,document_name) SELECT chunk_id,heading,text,document_name FROM chunks")
    con.execute("DELETE FROM questions_fts")
    con.execute("INSERT INTO questions_fts(question_id,question_text,professor,topic) SELECT question_id,question_text,professor,topic FROM questions")
    con.commit()
