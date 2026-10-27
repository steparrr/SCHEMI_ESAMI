from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path

import fitz
import openpyxl
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "scripts"))
from common import create_exam, read_json
from ingest import ingest
from formats import extract_pdf, extract_pptx, normalized_image, spreadsheet_blocks
from local_app import App, make_server


class FormatTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.base = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def test_excel_materials_preserve_sheets_cells_and_uncached_formulas(self):
        path = self.base / "materiale.xlsx"
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Criteri"
        worksheet.append(["Parametro", "Valore"])
        worksheet.append(["Alfa", "1,5"])
        worksheet.append(["Somma", "=1+2"])
        workbook.save(path)
        text = "\n".join(block["text"] for block in spreadsheet_blocks(path))
        self.assertIn("Criteri", text)
        self.assertIn("1,5", text)
        self.assertIn("=1+2 [risultato: non disponibile nel file]", text)

    def test_csv_semicolon_keeps_decimal_commas_and_quoted_cells(self):
        path = self.base / "domande.csv"
        path.write_text('Argomento;Domanda\nAlfa;"Valore 1,5; e criterio"\n')
        rows = spreadsheet_blocks(path)
        self.assertIn("Alfa | Valore 1,5; e criterio", rows[1]["text"])

    def test_pptx_reads_slide_text_notes_tables_and_images(self):
        png = self.base / "figura.png"
        Image.new("RGB", (600, 400), "#7297ad").save(png)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Concetto Alfa"
        slide.placeholders[1].text = "Definizione precisa e spiegazione documentata della figura. " * 4
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Parametro"
        table.cell(1, 0).text = "Alfa"
        table.cell(1, 1).text = "1,5"
        slide.shapes.add_picture(str(png), Inches(5), Inches(1), width=Inches(2))
        slide.notes_slide.notes_text_frame.text = "Nota del docente da conservare."
        path = self.base / "slide.pptx"
        deck.save(path)
        blocks, images, pages = extract_pptx(path)
        self.assertEqual(pages, 1)
        self.assertIn("Nota del docente da conservare", blocks[0]["text"])
        self.assertIn("Alfa | 1,5", blocks[0]["text"])
        self.assertEqual(len(images), 1)

    def test_heic_is_decoded_and_normalized_for_ocr_and_api(self):
        from pillow_heif import register_heif_opener
        register_heif_opener()
        path = self.base / "foto.heic"
        Image.new("RGB", (320, 250), "#748ca1").save(path, format="HEIF")
        data = normalized_image(path)
        with Image.open(io.BytesIO(data)) as picture:
            self.assertEqual(picture.format, "PNG")
            self.assertEqual(picture.size, (320, 250))

    @unittest.skipUnless(shutil.which("tesseract"), "OCR locale non disponibile sul sistema di collaudo")
    def test_scanned_pdf_is_ocr_read_and_keeps_page_provenance(self):
        image = Image.new("RGB", (1600, 700), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
        draw.text((90, 90), "ALFA CONCETTI FONDAMENTALI", font=font, fill="black")
        draw.text((90, 180), "Definizione e criteri precisi.", font=font, fill="black")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        path = self.base / "scansione.pdf"
        with fitz.open() as document:
            page = document.new_page(width=800, height=350)
            page.insert_image(page.rect, stream=buffer.getvalue())
            document.save(path)
        blocks, images, pages = extract_pdf(path)
        self.assertEqual(pages, 1)
        self.assertEqual(blocks[0]["page_start"], 1)
        self.assertIn("ALFA", blocks[0]["text"])
        self.assertIn("criteri", blocks[0]["text"])
        self.assertTrue(images)

    def test_unsupported_file_stops_import_instead_of_ignoring_it(self):
        root = create_exam("Prova formati", self.base)
        (root / "input/sbobine/vecchio.doc").write_bytes(b"vecchio file")
        with self.assertRaises(ValueError):
            ingest(root, lambda _: None)
        self.assertIn("input/sbobine/vecchio.doc", read_json(root / "output/reports/ingest.json")["unsupported"])


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.base = Path(self.temporary.name)
        shutil.copytree(PROJECT / "config", self.base / "config")
        shutil.copytree(PROJECT / "ui", self.base / "ui")
        self.app = App(self.base)
        self.server = make_server(self.app)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()
        self.temporary.cleanup()

    def request(self, path, body=None, token=True, headers=None):
        merged = {"X-Schemi-Token": self.app.token} if token else {}
        merged.update(headers or {})
        request = urllib.request.Request(self.url + path, headers=merged, data=json.dumps(body).encode() if body is not None else None)
        with urllib.request.urlopen(request, timeout=5) as response:
            return json.load(response)

    def test_requests_require_session_token_and_reject_other_origins(self):
        with self.assertRaises(urllib.error.HTTPError) as missing:
            self.request("/api/state", token=False)
        self.assertEqual(missing.exception.code, 403)
        with self.assertRaises(urllib.error.HTTPError) as external:
            self.request("/api/state", headers={"Origin": "https://example.invalid"})
        self.assertEqual(external.exception.code, 403)
        self.assertEqual(self.request("/api/state")["exams"], [])

    def test_create_fast_exam_and_paste_questions(self):
        result = self.request("/api/save", {"name": "Materia di prova", "mode": "veloce", "max_pages": 8, "questions_text": "Descrivi alfa\nDescrivi beta", "active": False})
        detail = self.request("/api/exam?slug=" + result["slug"])
        self.assertEqual(detail["exam"]["max_pagine"], 8)
        self.assertFalse(detail["exam"]["attivo"])
        self.assertIn("Descrivi alfa", detail["questions_text"])
        self.assertTrue(detail["pending_changes"])

    def test_upload_duplicate_is_reused_and_changed_same_name_is_preserved(self):
        slug = self.app.save_exam({"name": "Prova upload"})
        first = self.app.upload(slug, "materiali", "appunti.txt", io.BytesIO(b"Alfa definizione"), 16)
        second = self.app.upload(slug, "materiali", "appunti.txt", io.BytesIO(b"Alfa definizione"), 16)
        third = self.app.upload(slug, "materiali", "appunti.txt", io.BytesIO(b"Beta definizione"), 16)
        self.assertEqual(first["name"], "appunti.txt")
        self.assertTrue(second["duplicate"])
        self.assertEqual(third["name"], "appunti (2).txt")

    def test_upload_path_traversal_is_rejected(self):
        slug = self.app.save_exam({"name": "Prova percorso"})
        with self.assertRaises(ValueError):
            self.app.upload(slug, "materiali", "../esame.json", io.BytesIO(b"{}"), 2)
        self.assertEqual(read_json(self.app.root(slug) / "esame.json")["nome"], "Prova percorso")

    def test_default_model_changes_preserve_existing_exam_profile(self):
        result = self.request("/api/save", {"name": "Prima materia"})
        profile = self.app.defaults()
        original = profile["model"]
        profile.update(model="modello-futuro-di-prova", reasoning_effort="")
        self.request("/api/settings", profile)
        self.assertEqual(self.app.detail(result["slug"])["exam"]["pipeline"]["model"], original)
        second = self.request("/api/save", {"name": "Seconda materia"})
        self.assertEqual(self.app.detail(second["slug"])["exam"]["pipeline"]["model"], "modello-futuro-di-prova")


if __name__ == "__main__":
    unittest.main()
